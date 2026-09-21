# -*- coding: utf-8 -*-
import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    return p

def format_inline(paragraph, text, default_color=RGBColor(51, 65, 85), default_size=Pt(10.5)):
    # Split by bold tags **text** and code `text`
    parts = re.split(r'(\*\*.*?\*\*|\`.*?\`|\[.*?\]\(.*?\))', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = default_size
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(15, 118, 110)
        elif part.startswith('[') and '](' in part:
            match = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if match:
                link_text, link_url = match.groups()
                run = paragraph.add_run(f"{link_text} ({link_url})")
                run.font.name = 'Calibri'
                run.font.size = default_size
                run.font.color.rgb = RGBColor(5, 150, 105)
                run.underline = True
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Calibri'
            run.font.size = default_size
            run.font.color.rgb = default_color

def create_docx(readme_path, output_docx_path):
    doc = Document()
    
    # Page Setup - 0.8 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    with open(readme_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        rows_data = []
        for l in t_lines:
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', l):
                continue
            cells = [c.strip() for c in l.strip().strip('|').split('|')]
            if cells and any(cells):
                rows_data.append(cells)
        
        if not rows_data:
            return
            
        num_cols = max(len(r) for r in rows_data)
        # Pad shorter rows
        for r in rows_data:
            while len(r) < num_cols:
                r.append('')
                
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        
        for row_idx, row_data in enumerate(rows_data):
            row = table.rows[row_idx]
            is_header = (row_idx == 0)
            for col_idx, cell_value in enumerate(row_data):
                cell = row.cells[col_idx]
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                if is_header:
                    set_cell_background(cell, '059669') # Emerald accent
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run(cell_value)
                    run.bold = True
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    bg_col = 'F8FAFC' if row_idx % 2 == 1 else 'FFFFFF'
                    set_cell_background(cell, bg_col)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    format_inline(p, cell_value, default_size=Pt(9.0))
        
        # Add small spacing after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)

    def flush_code_block(c_lines):
        if not c_lines:
            return
        code_text = "".join(c_lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        set_cell_background(cell, 'F1F5F9')
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        for line in c_lines:
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block(code_block_lines)
                code_block_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if '|' in line and (line.strip().startswith('|') or line.strip().endswith('|') or '-|-' in line):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal Divider
        if stripped in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, '', space_before=14, space_after=6)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = add_styled_paragraph(doc, '', space_before=12, space_after=4)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(5, 150, 105) # Emerald
        elif stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = add_styled_paragraph(doc, '', space_before=10, space_after=3)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = add_styled_paragraph(doc, '', space_before=8, space_after=2)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(71, 85, 105)
        # Blockquotes
        elif stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, '', space_before=4, space_after=4)
            p.paragraph_format.left_indent = Inches(0.25)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="059669"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            format_inline(p, quote_text, default_color=RGBColor(71, 85, 105))
        # Bullet list
        elif stripped.startswith('- ') or stripped.startswith('* '):
            b_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, '', space_before=2, space_after=3)
            p.paragraph_format.left_indent = Inches(0.25)
            # Add custom bullet dot
            run_bullet = p.add_run("\u2022  ")
            run_bullet.bold = True
            run_bullet.font.color.rgb = RGBColor(5, 150, 105)
            format_inline(p, b_text)
        # Numbered list
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^(\d+\.)\s+(.*)', stripped)
            prefix, n_text = match.groups()
            p = add_styled_paragraph(doc, '', space_before=2, space_after=3)
            p.paragraph_format.left_indent = Inches(0.25)
            run_num = p.add_run(f"{prefix} ")
            run_num.bold = True
            run_num.font.color.rgb = RGBColor(5, 150, 105)
            format_inline(p, n_text)
        else:
            p = add_styled_paragraph(doc, '', space_before=0, space_after=4)
            format_inline(p, stripped)

        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code_block(code_block_lines)

    doc.save(output_docx_path)
    print(f"Successfully generated DOCX: {output_docx_path}")

if __name__ == "__main__":
    readme = os.path.abspath("README.md")
    output_docx = os.path.abspath("myHR_by_Swaniki_README.docx")
    create_docx(readme, output_docx)

